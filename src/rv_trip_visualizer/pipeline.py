"""Ties parsing, merging, integrity checking, and output writing together.
Used by the UI; also usable standalone for scripting.
"""
from __future__ import annotations

import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import integrity, itinerary_page, map_page

OUTPUT_KINDS = ["data_json", "xlsx", "itinerary_html", "map_html"]


class NotClean(Exception):
    """Raised by build() when the integrity check isn't clean. Carries the
    report so the caller can show the same findings a Check Data run would.
    """

    def __init__(self, report: integrity.Report):
        self.report = report
        super().__init__(
            f"{len(report.findings)} integrity finding(s) - fix the source "
            "data and re-check before building."
        )


def check(paths: list[str]) -> integrity.Report:
    return integrity.check(paths)


def build(
    paths: list[str],
    out_dir: str,
    base_name: str,
    itinerary_title: str,
    map_title: str,
) -> dict:
    """Runs its own fresh integrity check and refuses to write anything
    unless it comes back completely clean - this is where the "no output
    until the data is right" rule is actually enforced, independent of
    whatever the UI last displayed.
    """
    report = integrity.check(paths)
    if not report.is_clean:
        raise NotClean(report)

    stays = report.merged_stays
    os.makedirs(out_dir, exist_ok=True)

    paths_out = {
        "data_json": os.path.join(out_dir, f"{base_name}-data.json"),
        "xlsx": os.path.join(out_dir, f"{base_name}.xlsx"),
        "itinerary_html": os.path.join(out_dir, f"{base_name}-Itinerary.html"),
        "map_html": os.path.join(out_dir, f"{base_name}-Map.html"),
    }

    with open(paths_out["data_json"], "w", encoding="utf-8") as f:
        json.dump(stays, f, indent=1, default=str)

    _write_itinerary_xlsx(stays, paths_out["xlsx"])
    itinerary_page.render(stays, paths_out["itinerary_html"], itinerary_title)
    map_page.render(stays, paths_out["map_html"], map_title)

    return {"stays": stays, "log": report.log, "outputs": paths_out}


def _write_itinerary_xlsx(stays: list[dict], out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Itinerary"
    headers = ["Arrival", "Departure", "Nights", "City", "State", "Year", "Miles", "Lat", "Lon"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(1, col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5233")
    border = Border(bottom=Side(style="thin", color="D8D5C8"))
    for s in stays:
        ws.append([
            s["arrival"], s["departure"], s["nights"], s["city"], s["state"],
            s["arrival"].year, s["miles"], round(s["lat"], 5), round(s["lon"], 5),
        ])
        for col in range(1, len(headers) + 1):
            ws.cell(ws.max_row, col).border = border
    widths = {1: 12, 2: 12, 3: 9, 4: 22, 5: 12, 6: 8, 7: 9, 8: 10, 9: 10}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("Notes")
    notes.append(["RV Trip Visualizer - build notes"])
    notes["A1"].font = Font(bold=True, size=13)
    notes.append([])
    notes.append(["Source: the selected export files, merged chronologically, overnight stays only (nights >= 1)."])
    notes.append(["Where two exports covered the same dates (the plan changed mid-trip), the later export was treated as authoritative."])
    notes.append(["Miles = the source tool's own Miles column, with any 0-night waypoint stops (gas, lunch, etc.) folded forward into the next overnight stay, since those waypoints aren't shown as their own rows here."])
    notes.column_dimensions["A"].width = 100
    for row in notes.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out_path)

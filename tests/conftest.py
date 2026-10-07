"""Shared test helpers - build small, in-memory-shaped xlsx workbooks that
mimic a real export's layout (a title row, a "Start Date:" label row, a
blank row, then the header row and data) without any checked-in binary
fixtures or real trip data.
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

ALL_HEADERS = [
    "Stop Name", "Miles", "Total", "Estimated Travel Time", "Arrival Day",
    "Arrival Date", "Nights", "Departure Day", "Departure Date", "Comments",
    "Reservation Number", "Features", "Location", "Url", "Phone", "Email",
    "Latitude", "Longitude", "Camping Cost", "Meals Cost",
]


@pytest.fixture
def make_workbook(tmp_path):
    def _make(
        filename: str,
        headers: list[str],
        rows: list[dict],
        start_date_label: str = "Start Date: Thursday, June 27, 2019",
        sheet_name: str = "Trip Summary",
    ) -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.cell(1, 1, "2019 [06-06] Test Trip")
        ws.cell(2, 1, start_date_label)
        header_row = 4
        for c, h in enumerate(headers, start=1):
            ws.cell(header_row, c, h)
        for r, row in enumerate(rows, start=header_row + 1):
            for c, h in enumerate(headers, start=1):
                if h in row:
                    ws.cell(r, c, row[h])
        path = tmp_path / filename
        wb.save(path)
        return str(path)

    return _make

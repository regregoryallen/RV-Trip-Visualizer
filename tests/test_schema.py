import openpyxl

from rv_trip_visualizer import schema
from .conftest import ALL_HEADERS


def test_resolve_headers_finds_full_header_set(make_workbook):
    path = make_workbook("full.xlsx", ALL_HEADERS, [
        {"Stop Name": "Austin, TX", "Miles": 10, "Total": 10, "Arrival Date": "6/27/19",
         "Nights": 1, "Departure Date": "6/28/19", "Location": "Austin, TX",
         "Latitude": 30.3, "Longitude": -97.7},
    ])
    wb = openpyxl.load_workbook(path, data_only=True)
    hm = schema.resolve_headers(wb)
    assert hm is not None
    assert hm.sheet_name == "Trip Summary"
    assert hm.header_row == 4
    assert set(hm.columns) == set(schema.REQUIRED_HEADERS)


def test_resolve_headers_tolerates_reordered_and_extra_columns(make_workbook):
    shuffled = list(reversed(ALL_HEADERS)) + ["Some Extra Column"]
    path = make_workbook("shuffled.xlsx", shuffled, [
        {"Stop Name": "Austin, TX", "Miles": 10, "Total": 10, "Arrival Date": "6/27/19",
         "Nights": 1, "Departure Date": "6/28/19", "Location": "Austin, TX",
         "Latitude": 30.3, "Longitude": -97.7},
    ])
    wb = openpyxl.load_workbook(path, data_only=True)
    hm = schema.resolve_headers(wb)
    assert hm is not None
    assert set(hm.columns) == set(schema.REQUIRED_HEADERS)


def test_resolve_headers_returns_none_when_a_required_header_is_missing(make_workbook):
    missing_total = [h for h in ALL_HEADERS if h != "Total"]
    path = make_workbook("missing_total.xlsx", missing_total, [
        {"Stop Name": "Austin, TX", "Miles": 10, "Arrival Date": "6/27/19",
         "Nights": 1, "Departure Date": "6/28/19", "Location": "Austin, TX",
         "Latitude": 30.3, "Longitude": -97.7},
    ])
    wb = openpyxl.load_workbook(path, data_only=True)
    assert schema.resolve_headers(wb) is None
    missing, closest_sheet = schema.missing_headers(wb)
    assert missing == ["Total"]
    assert closest_sheet == "Trip Summary"
